"""KBR H2ACT OSBL equipment sizing workbook -- 12 / 24 / 80 / 160 kTPA H2.

Builds output/KBR_OSBL_Equipment_Sizing.xlsx (git-ignored; regenerate with
`python tools/kbr_osbl_workbook.py`, then recalculate with LibreOffice or open
in Excel). Every result cell is a live formula driven by:

  Inputs    -- user-editable design assumptions (yellow), one named range each
  KBR_Data  -- KBR figures with document citations (white) + scaling rules

Sheet layout follows mdlguideline.md (Cover / Guide / Inputs / reference /
calculation / Summary). Method and assumptions mirror
tcoedatabase/KBR_OSBL_Sizing_12ktpa.md and tools/kbr_osbl_sizing_12ktpa.py.

KBR sources (Licensor/kbr/, Rev 0, Dec 2025):
  I.A Process Design Basis (Gentari-PR-GEN-PDB-001)
  I.C Heat & Material Balance, 12 kTPA NG mode (Gentari-PR-GEN-HMB-001)
  I.E Feed & Product Summary, all capacities (Gentari-PR-GEN-F&F-001)
  I.F Utility Summary ISBL, 12 kTPA only (Gentari-PR-GEN-BLS-001)
  I.H Equipment List OSBL (Gentari-PR-GEN-LST-0002)
  hub H2ACT Technical Information Package (kbr-johor-hub.md)
"""
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.protection import Protection
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

OUT = Path(__file__).resolve().parents[1] / "output" / "KBR_OSBL_Equipment_Sizing.xlsx"
VERSION, DATE = "v1.0", "2026-09-30"

# mdlguideline.md §5-6
F_INPUT, F_CALC, F_REF, F_HEAD, F_WARN, F_EXTRAP, F_TBD = (
    "FFF2CC", "F2F2F2", "FFFFFF", "1F3864", "FFE699", "FCE4D6", "F8CBAD")
FONT = "Calibri"
NUM = '[>=100]#,##0;[>=10]#,##0.0;0.00'
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def fill(c):
    return PatternFill(start_color=c, end_color=c, fill_type="solid")


def font(size=11, bold=False, color="000000", italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=underline)


def style(cell, bg=None, bold=False, color="000000", fmt=None, wrap=True, box=True, size=11,
          align="left"):
    cell.font = font(size=size, bold=bold, color=color)
    if bg:
        cell.fill = fill(bg)
    if fmt:
        cell.number_format = fmt
    cell.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal=align)
    if box:
        cell.border = BOX


def header_row(ws, row, labels, widths=None):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=lab)
        style(c, bg=F_HEAD, bold=True, color="FFFFFF")
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = font(size=18, bold=True, color=F_HEAD)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = font(size=11, italic=True, color="595959")
    if ws.title != "Cover":
        c = ws["A3"]
        c.value = "◀ Back to Cover"
        c.hyperlink = Hyperlink(ref="A3", location="'Cover'!A1")
        c.font = font(color="0563C1", underline="single")


def add_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


# ---------------------------------------------------------------------------
# Inputs: (name, label, value, unit, source/assumption, validation (min, max))
# ---------------------------------------------------------------------------
INPUTS = [
    ("Design margins & philosophy", None),
    ("Margin", "Design margin on KBR maximum flows", 1.10, "×", "A1 ASSUMPTION: +10% on KBR maximum flows (all items except emergency power)", (1, 2)),
    ("StorageHours", "Storage autonomy for tanks", 24, "h", "A2 ASSUMPTION: extends the owner's 24 h diesel basis to every tank", (1, 720)),
    ("TankFill", "Net ÷ nominal tank volume (heel + freeboard)", 0.90, "–", "A2 ASSUMPTION", (0.5, 1)),
    ("Emergency power (owner method)", None),
    ("PowerMult", "Load allowance multiplier on KBR max ISBL power (NOT redundancy)", 2.0, "×", "A4 OWNER INSTRUCTION: ×2 covers loads I.F Note 7 excludes (CT pumps/fans, N₂ generator, lighting, buildings, instruments). Redundancy (e.g. N+1 gensets) would be a separate choice; I.H lists 1 package", (1, 5)),
    ("BackupHours", "Emergency backup duration", 24, "h", "A5 OWNER INSTRUCTION", (1, 168)),
    ("GenEff", "Genset (engine) efficiency", 0.50, "–", "A5 OWNER INSTRUCTION (check vendor heat rate; 40% → larger tank)", (0.2, 0.6)),
    ("PowerFactor", "Genset rating power factor", 0.80, "–", "Standard genset rating basis (ISO 8528 rates at 0.8 pf)", (0.5, 1)),
    ("DieselLHV", "Diesel LHV", 43.0, "MJ/kg", "IPCC 2006 Guidelines Vol.2 Ch.1 Table 1.2 (gas/diesel oil NCV 43.0 TJ/Gg)", (35, 46)),
    ("DieselDensity", "Diesel density", 820, "kg/m³", "EN 590: 820–845 kg/m³ @15 °C; low end → larger tank", (780, 880)),
    ("Water systems", None),
    ("COC", "Cooling tower cycles of concentration", 5, "–", "A6 ASSUMPTION (water-treatment vendor to confirm)", (1.5, 15)),
    ("DeminRecovery", "Demin plant recovery (product ÷ feed)", 0.80, "–", "A7 ASSUMPTION", (0.5, 1)),
    ("Headcount", "Persons on site (all capacities)", 30, "persons", "A8 ASSUMPTION — OWNER TO CONFIRM (likely higher at 80/160 kTPA)", (1, 1000)),
    ("PotableLPD", "Potable water per person", 100, "L/person/day", "WHO/UN 50–100 L/person/day (upper end)", (20, 500)),
    ("DomesticPeak", "Domestic peak factor", 3, "×", "A8 ASSUMPTION", (1, 10)),
    ("Showers", "Simultaneous safety showers", 2, "no.", "A8 ASSUMPTION", (1, 10)),
    ("ShowerFlow", "Safety shower flow", 75.7, "L/min", "ANSI/ISEA Z358.1: 20 gpm", (50, 200)),
    ("ShowerMin", "Safety shower duration", 15, "min", "ANSI/ISEA Z358.1", (5, 60)),
    ("Cp", "Water specific heat", 4.186, "kJ/kg·K", "Standard property", (4, 4.3)),
    ("Hfg", "Latent heat of water @30 °C", 2430, "kJ/kg", "Steam tables (IAPWS-IF97)", (2300, 2500)),
    ("CWPumpDP", "CW pump differential pressure (power cross-check)", 5.5, "bar", "ASSUMPTION (cross-check only)", (1, 15)),
    ("PumpEff", "Pump efficiency (power cross-check)", 0.70, "–", "ASSUMPTION (cross-check only)", (0.3, 0.9)),
    ("Air & nitrogen", None),
    ("DryerPurge", "Heatless IA dryer purge fraction", 0.15, "–", "A9 ASSUMPTION", (0, 0.3)),
    ("CompDisch", "Air compressor discharge pressure", 8.5, "barg", "A9 ASSUMPTION: 1 bar above 7.5 barg header (I.F)", (7.5, 15)),
    ("IsoEff", "Compressor isothermal efficiency (power cross-check)", 0.65, "–", "ASSUMPTION (cross-check only)", (0.4, 0.9)),
    ("IAHoldMin", "Instrument air receiver hold-up", 15, "min", "A10 ASSUMPTION", (1, 60)),
    ("IAMinP", "Minimum IA pressure at end of hold-up", 4.0, "barg", "A10 ASSUMPTION", (1, 7)),
    ("WetHoldMin", "Wet air receiver hold-up (of compressor output)", 2, "min", "A10 ASSUMPTION", (0.5, 30)),
    ("N2StartupH", "Duration of start-up N₂ peak", 24, "h", "A11 ASSUMPTION — KBR says peak TBC (I.F Notes 5–6)", (1, 168)),
    ("LN2Rho", "Liquid N₂ density @ normal boiling point", 806.6, "kg/m³", "Standard property", (800, 810)),
    ("N2GasRho", "N₂ gas density @0 °C, 1 atm", 1.2506, "kg/Nm³", "Standard property", (1.2, 1.3)),
    ("Hydrogen export & general", None),
    ("PipeVmax", "Max H₂ export line velocity", 20, "m/s", "A12 ASSUMPTION", (5, 40)),
    ("NM3perKmol", "Molar volume @0 °C, 1 atm", 22.414, "Nm³/kmol", "Ideal gas (KBR Nm³ basis: 0 °C, 1 atm — I.A §12)", (22.4, 22.42)),
    ("PatmBar", "Atmospheric pressure", 1.01325, "bar", "Standard atmosphere", (0.9, 1.1)),
]

# ---------------------------------------------------------------------------
# KBR_Data, capacity table: key -> (label, unit, [12, 24, 80], source)
# 160 kTPA column = 80 kTPA × (160/80): ASSUMPTION, beyond KBR's published range
# ---------------------------------------------------------------------------
CAPS = [("12 kTPA", 12), ("24 kTPA", 24), ("80 kTPA", 80), ("160 kTPA", 160)]
CAP_COLS = {12: "C", 24: "D", 80: "E", 160: "F"}
CAP_TABLE = [
    ("cap", "H₂ capacity", "kTPA", [12, 24, 80], "I.E / I.H title block"),
    ("p_norm", "Electrical power, normal operating load (highest published firing mode)", "kW", [487, 870, 2760],
     "I.E Energy Efficiency — Electrical Power Import (12 & 80: 100% cracked-gas mode is highest; 24: only 100% NG published)"),
    ("ng_norm", "NG fuel, 100% NG mode", "kg/h", [417.9, 831.4, 2759.3], "I.E FUELS — NG Fuel"),
    ("nh3_max", "NH₃ feed, highest-feed firing mode (as 100% NH₃)", "kg/h", [10290, 18102, 68470],
     "I.E FEED — Ammonia (12 & 80: 100% cracked-gas mode; 24: only 100% NG published)"),
    ("h2", "H₂ product (as 100% H₂)", "kg/h", [1426, 2851, 9505], "I.E PRODUCT — Hydrogen"),
]

# KBR_Data, single values: name -> (label, value, unit, source)
KBR_SINGLE = [
    ("I.F Utility Summary anchors (published for 12 kTPA only)", None),
    ("KBR_PMax12", "Electrical power, maximum (12 kTPA)", 540, "kW", "I.F Electrical Power, Maximum column (Notes 7, 9)"),
    ("KBR_PNorm12", "Electrical power, normal, highest mode (12 kTPA)", 487, "kW", "I.F Electrical Power, normal 100% CF (= I.E)"),
    ("KBR_PEmerg12", "KBR emergency backup power estimate (12 kTPA)", 300, "kW", "I.F Note 9: ~0.3 MW for L.O. pumps, instrumentation, MOVs"),
    ("KBR_NGMax12", "NG fuel, maximum (12 kTPA)", 450, "kg/h", "I.F Natural Gas, Maximum (covers start-up, Note 10)"),
    ("KBR_NGNorm12", "NG fuel, normal 100% NG (12 kTPA)", 418, "kg/h", "I.F Natural Gas, normal 100% NG"),
    ("KBR_CWMax12", "Cooling water, maximum (12 kTPA)", 40, "t/h", "I.F C.W. Supply, Maximum (includes start-up coolers)"),
    ("KBR_DeminMax12", "Demin water, maximum (12 kTPA)", 270, "kg/h", "I.F Demin. Water, Maximum"),
    ("KBR_PAMax12", "Plant air, maximum (12 kTPA)", 250, "Nm³/h", "I.F Plant Air, Maximum (Note 8: reference-plant estimate)"),
    ("KBR_IAMax12", "Instrument air, maximum (12 kTPA)", 350, "Nm³/h", "I.F Instrument Air, Maximum (Note 8)"),
    ("KBR_N2Nor12", "Nitrogen, normal (12 kTPA)", 150, "Nm³/h", "I.F Nitrogen, normal (Note 5: reference-plant estimate)"),
    ("KBR_N2Max12", "Nitrogen, maximum intermittent (12 kTPA)", 500, "Nm³/h", "I.F Nitrogen, Maximum (Note 6: start-up; peak TBC)"),
    ("KBR per-t H₂ ratios (constant 12–80 kTPA → basis for linear scaling)", None),
    ("KBR_CWRatio", "CW consumption", 25, "t/t H₂", "hub §3.3 Utility Demands (same at all capacities)"),
    ("KBR_DNWRatio", "Demin water consumption", 0.1, "t/t H₂", "hub §3.3 Utility Demands (same at all capacities)"),
    ("KBR_BDRatio", "Steam blowdown effluent", 0.06, "t/t H₂", "hub §3.5 (< 0.06 t/t H₂, only liquid effluent)"),
    ("Conditions & properties", None),
    ("KBR_CWSupplyT", "CW supply temperature", 30, "°C", "I.A §6 / I.D §6 (I.F says 25 °C — below design wet bulb, see Guide)"),
    ("KBR_CWReturnT", "CW return temperature", 40, "°C", "I.A §6 / I.D §6"),
    ("KBR_CWSupplyP", "CW supply pressure at BL", 5.0, "barg", "I.F (I.A: 2.0, I.D: 4.0 — highest used)"),
    ("KBR_DeminP", "Demin water pressure at ISBL BL", 5.0, "barg", "I.F / I.A §7"),
    ("KBR_AirP", "Plant / instrument air header pressure", 7.5, "barg", "I.F"),
    ("KBR_N2P", "Nitrogen pressure", 7.0, "barg", "I.F"),
    ("KBR_NGP", "NG supply pressure", 7.5, "barg", "I.F (I.A §5: 7 barg)"),
    ("KBR_NGMW", "NG molecular weight", 16.95, "kg/kmol", "I.C stream 13"),
    ("KBR_NGLHV", "NG LHV", 49.56, "MJ/kg", "I.E Gas Heating Values"),
    ("KBR_H2P", "H₂ product pressure at BL", 27, "barg", "I.E / I.C stream 8 (I.A: 20 min; hub: 28)"),
    ("KBR_H2MW", "H₂ product molecular weight", 2.02, "kg/kmol", "I.C stream 8"),
    ("KBR_H2Rho", "H₂ product density @27 barg, 30 °C", 2.21, "kg/m³", "I.C stream 8"),
    ("KBR_NH3MW", "NH₃ molecular weight", 17.031, "kg/kmol", "Standard (I.C: 17.04)"),
    ("KBR_Turndown", "Plant minimum turndown", 0.40, "–", "I.A §9"),
    ("KBR_AmbT", "Ambient temperature (annual average)", 28, "°C", "I.A §4"),
    ("KBR_AmbRH", "Relative humidity", 80, "%", "I.A §4"),
]

# ASME B36.10M Schedule 40 (STD up to DN250) inside diameters, mm
PIPES = [("DN100 (4\")", 102.3), ("DN150 (6\")", 154.1), ("DN200 (8\")", 202.7), ("DN250 (10\")", 254.5),
         ("DN300 (12\")", 303.2), ("DN350 (14\")", 333.4), ("DN400 (16\")", 381.0), ("DN450 (18\")", 428.7),
         ("DN500 (20\")", 477.8), ("DN600 (24\")", 574.6)]

# ---------------------------------------------------------------------------
# Calculations on each capacity sheet: key -> (label, formula, unit, equation/ref)
# {key} -> cell of that key on the same sheet; {COL} -> KBR_Data capacity column
# ---------------------------------------------------------------------------
CALCS = [
    ("A. Capacity basis (KBR data for this capacity)", None),
    ("cap", "H₂ capacity", "=KBR_Data!{COL}$6", "kTPA", "KBR_Data"),
    ("scale", "Scale factor vs 12 kTPA anchor", "={cap}/KBR_Data!$C$6", "×", "cap ÷ 12"),
    ("p_norm", "Electrical power, normal (highest mode)", "=KBR_Data!{COL}$7", "kW", "I.E"),
    ("ng_norm", "NG fuel, 100% NG mode", "=KBR_Data!{COL}$8", "kg/h", "I.E"),
    ("nh3_max", "NH₃ feed, highest-feed mode", "=KBR_Data!{COL}$9", "kg/h", "I.E"),
    ("h2", "H₂ product", "=KBR_Data!{COL}$10", "kg/h", "I.E"),
    ("B. KBR utility maxima at this capacity", None),
    ("p_max", "Electrical power, maximum", "={p_norm}*KBR_PMax12/KBR_PNorm12", "kW", "P_max = P_norm × (540/487) — I.F max/normal ratio at 12 kTPA applied [scaling rule S1]"),
    ("ng_max", "NG fuel, maximum", "={ng_norm}*KBR_NGMax12/KBR_NGNorm12", "kg/h", "NG_max = NG_norm × (450/418) — I.F ratio at 12 kTPA [S1]"),
    ("cw_max", "Cooling water, maximum", "=KBR_CWMax12*{scale}", "t/h", "40 t/h × scale — linear, per hub §3.3 constant 25 t/t H₂ [S2]"),
    ("demin_max", "Demin water, maximum", "=KBR_DeminMax12*{scale}", "kg/h", "270 kg/h × scale — linear, per hub §3.3 constant 0.1 t/t H₂ [S2]"),
    ("pa_max", "Plant air, maximum", "=KBR_PAMax12*{scale}", "Nm³/h", "250 × scale [S3 ASSUMPTION: linear]"),
    ("ia_max", "Instrument air, maximum", "=KBR_IAMax12*{scale}", "Nm³/h", "350 × scale [S3]"),
    ("n2_nor", "Nitrogen, normal", "=KBR_N2Nor12*{scale}", "Nm³/h", "150 × scale [S3]"),
    ("n2_max", "Nitrogen, maximum (start-up)", "=KBR_N2Max12*{scale}", "Nm³/h", "500 × scale [S3]"),
    ("C. Unit 101 — Emergency power", None),
    ("gen_kw", "Genset rating", "={p_max}*PowerMult", "kWe", "P_gen = P_max × PowerMult (owner method)"),
    ("gen_kva", "Genset apparent power", "={gen_kw}/PowerFactor", "kVA", "S = P ÷ pf"),
    ("e_el", "Electrical energy over backup period", "={gen_kw}*BackupHours", "kWh", "E = P_gen × t"),
    ("fuel_mj", "Fuel energy required", "={e_el}/GenEff*3.6", "MJ", "E_fuel = E ÷ η × 3.6 MJ/kWh"),
    ("diesel_kg", "Diesel mass", "={fuel_mj}/DieselLHV", "kg", "m = E_fuel ÷ LHV"),
    ("burn_lph", "Diesel burn rate", "={diesel_kg}/BackupHours/DieselDensity*1000", "L/h", "m ÷ t ÷ ρ"),
    ("diesel_net", "Diesel tank, net", "={diesel_kg}/DieselDensity", "m³", "V = m ÷ ρ"),
    ("diesel_nom", "Diesel tank, nominal", "={diesel_net}/TankFill", "m³", "V_nom = V_net ÷ fill"),
    ("sens40", "Sensitivity: net tank at 40% efficiency", "={e_el}/0.4*3.6/DieselLHV/DieselDensity", "m³", "Same with η = 40%"),
    ("sens_n9", "Sensitivity: net tank on KBR Note 9 basis (0.3 MW × scale)", "=KBR_PEmerg12*{scale}*BackupHours/GenEff*3.6/DieselLHV/DieselDensity", "m³", "I.F Note 9, scaled linearly"),
    ("D. Unit 107 — Cooling water", None),
    ("cw_des", "CW circulation, design", "={cw_max}*Margin", "m³/h", "CW_max × margin"),
    ("cw_each", "Per pump (3 × 50%) / per cell (2 cells)", "={cw_des}/2", "m³/h", "CW_des ÷ 2"),
    ("cw_duty", "CW heat duty", "={cw_des}*1000*Cp*(KBR_CWReturnT-KBR_CWSupplyT)/3600", "kW", "Q = ṁ × cp × ΔT"),
    ("evap", "CT evaporation", "={cw_duty}*3600/Hfg", "kg/h", "E = Q ÷ h_fg (all heat by evaporation, conservative)"),
    ("ct_bd", "CT blowdown", "={evap}/(COC-1)", "kg/h", "B = E ÷ (COC − 1)"),
    ("ct_mu", "CT make-up", "={evap}+{ct_bd}", "kg/h", "M = E + B (drift neglected)"),
    ("cw_pump_kw", "CW pump shaft power (cross-check)", "={cw_des}/3600*CWPumpDP*100000/PumpEff/1000", "kW", "P = Q × ΔP ÷ η"),
    ("E. Units 103 / 104 — Demin & polished water", None),
    ("demin_des", "Demin / polished water, design", "={demin_max}*Margin/1000", "m³/h", "Demin_max × margin (≈1,000 kg/m³)"),
    ("demin_net", "Demin / polished tank, net", "={demin_des}*StorageHours", "m³", "Q × storage h"),
    ("demin_nom", "Demin / polished tank, nominal", "={demin_net}/TankFill", "m³", "÷ fill"),
    ("demin_feed", "Raw water to demin package", "={demin_des}/DeminRecovery", "m³/h", "Q ÷ recovery"),
    ("demin_waste", "Demin regeneration / reject waste", "={demin_feed}-{demin_des}", "m³/h", "Feed − product"),
    ("neut_net", "Neutralisation tank, net", "={demin_waste}*StorageHours", "m³", "Waste × storage h"),
    ("neut_nom", "Neutralisation tank, nominal", "={neut_net}/TankFill", "m³", "÷ fill"),
    ("F. Unit 105 — Potable water", None),
    ("pot_day", "Potable demand", "=Headcount*PotableLPD/1000", "m³/d", "Persons × L/p/d"),
    ("pot_avg", "Potable demand, average", "={pot_day}/24", "m³/h", "÷ 24"),
    ("shower_vol", "Safety shower reserve", "=Showers*ShowerFlow*ShowerMin/1000", "m³", "n × Q × t"),
    ("pot_net", "Potable tank, net", "={pot_avg}*StorageHours+{shower_vol}", "m³", "Domestic × storage h + showers"),
    ("pot_nom", "Potable tank, nominal", "={pot_net}/TankFill", "m³", "÷ fill"),
    ("pot_pump", "Potable pump rate", "=Showers*ShowerFlow*60/1000+DomesticPeak*{pot_avg}", "m³/h", "Showers + peak domestic"),
    ("G. Unit 102 — Raw water", None),
    ("raw", "Raw water, continuous", "={ct_mu}/1000+{demin_feed}+{pot_avg}", "m³/h", "CT make-up + demin feed + potable"),
    ("raw_net", "Raw water tank, net (excl. fire reserve)", "={raw}*StorageHours", "m³", "Q × storage h"),
    ("raw_nom", "Raw water tank, nominal (excl. fire reserve)", "={raw_net}/TankFill", "m³", "÷ fill"),
    ("H. Unit 108 — Flare & fuel gas", None),
    ("flare_kg", "Flare load, proxy (full cracked gas)", "={nh3_max}*Margin", "kg/h", "A14 PROXY: NH₃ feed mass, fully cracked × margin — replace with relief study"),
    ("flare_nm3", "Flare load, volumetric", "={nh3_max}/KBR_NH3MW*2*NM3perKmol*Margin", "Nm³/h", "2 NH₃ → N₂ + 3 H₂ : 2 kmol out per kmol NH₃"),
    ("flare_mw", "Cracked gas molecular weight", "=KBR_NH3MW/2", "kg/kmol", "H₂/N₂ 3:1"),
    ("ng_des", "Fuel gas, design", "={ng_max}*Margin", "kg/h", "NG_max × margin"),
    ("ng_nm3", "Fuel gas, volumetric", "={ng_des}/KBR_NGMW*NM3perKmol", "Nm³/h", "ṁ ÷ MW × 22.414"),
    ("ng_duty", "Fuel gas firing duty (LHV)", "={ng_des}*KBR_NGLHV/3600", "MW", "ṁ × LHV"),
    ("I. Unit 109 — Air", None),
    ("ia_des", "Instrument air, design (dry product)", "={ia_max}*Margin", "Nm³/h", "IA_max × margin"),
    ("dryer_in", "IA dryer inlet", "={ia_des}/(1-DryerPurge)", "Nm³/h", "÷ (1 − purge)"),
    ("comp_tot", "Air compressors, total", "={pa_max}*Margin+{dryer_in}", "Nm³/h", "PA × margin + dryer inlet"),
    ("comp_each", "Per compressor (3 × 50%)", "={comp_tot}/2", "Nm³/h", "÷ 2 duty"),
    ("comp_kw", "Compressor shaft power (cross-check)", "={comp_tot}*PatmBar*100000*LN((CompDisch+PatmBar)/PatmBar)/3600000/IsoEff", "kW", "W = p₁V₁ ln(p₂/p₁) ÷ η_iso"),
    ("ia_rcv", "IA receiver volume", "={ia_max}*IAHoldMin/60*PatmBar/(KBR_AirP-IAMinP)", "m³", "V = Q × t × p_atm ÷ Δp"),
    ("wet_rcv", "Wet air receiver volume", "={comp_tot}*WetHoldMin/60*PatmBar/(KBR_AirP-IAMinP)", "m³", "V = Q × t × p_atm ÷ Δp"),
    ("J. Unit 110 — Nitrogen", None),
    ("n2_gen", "N₂ generation, design", "={n2_nor}*Margin", "Nm³/h", "N2_nor × margin"),
    ("ln2_net", "Liquid N₂ storage, net", "=MAX({n2_max}-{n2_gen},0)*N2StartupH/(LN2Rho/N2GasRho)", "m³", "(peak − generator) × t ÷ 645 Nm³/m³"),
    ("ln2_nom", "Liquid N₂ storage, nominal", "={ln2_net}/TankFill", "m³", "÷ fill"),
    ("vap_each", "Ambient vaporiser (each, 2 × 100%)", "={n2_max}*Margin", "Nm³/h", "N2_max × margin"),
    ("K. Unit 112 — Waste water", None),
    ("steam_bd", "Steam blowdown effluent", "=KBR_BDRatio*{h2}", "kg/h", "0.06 t/t H₂ × H₂ (hub §3.5)"),
    ("effluent", "Continuous effluent, total", "={steam_bd}/1000+{ct_bd}/1000+{demin_waste}", "m³/h", "Steam BD + CT BD + demin waste (storm water excluded)"),
    ("L. Unit 113 — Hydrogen export", None),
    ("h2_des", "H₂ metering, design", "={h2}*Margin", "kg/h", "H₂ × margin"),
    ("h2_nm3", "H₂ metering, volumetric", "={h2_des}/KBR_H2MW*NM3perKmol", "Nm³/h", "ṁ ÷ MW × 22.414"),
    ("h2_am3", "H₂ actual flow @27 barg, 30 °C", "={h2_des}/KBR_H2Rho", "Am³/h", "ṁ ÷ ρ (I.C)"),
    ("h2_min", "H₂ minimum flow (turndown)", "={h2}*KBR_Turndown", "kg/h", "H₂ × 40%"),
    ("h2_td", "Required meter turndown", "={h2_des}/{h2_min}", ": 1", "Design ÷ minimum"),
    ("pipe_id", "Minimum export line ID", "=SQRT(4*({h2_am3}/3600)/PI()/PipeVmax)*1000", "mm", "d = √(4Q ÷ πv)"),
    ("pipe_idx", "Pipe table row (smallest Sch 40 ID ≥ minimum)", "=COUNTIF(PipeID,\"<\"&{pipe_id})+1", "–", "Lookup into KBR_Data pipe table"),
    ("pipe_dn", "Indicative export line size", "=IFERROR(INDEX(PipeDN,{pipe_idx}),\"> DN600 — check\")", "–", "ASME B36.10M Sch 40"),
    ("pipe_v", "Velocity in selected line", "=IFERROR({h2_am3}/3600/(PI()/4*(INDEX(PipeID,{pipe_idx})/1000)^2),\"n/a\")", "m/s", "v = Q ÷ A"),
    ("M. Power cross-check (×2 allowance vs estimated OSBL rotating load)", None),
    ("allow", "Allowance added by the multiplier", "={gen_kw}-{p_max}", "kW", "P_gen − P_max"),
    ("osbl_est", "Estimated OSBL load (compressors + CW pumps only)", "={comp_kw}+{cw_pump_kw}", "kW", "N₂ gen, CT fans, lighting, buildings not included"),
]

# Equipment list: (unit, tag, item, qty, primary formula, unit, secondary formula, unit, basis, ref, confidence)
TBD = "TBD"
EQUIP = [
    ("101 Emergency Power / Sub-Station", "111-L", "Emergency Power Package", "1", "={gen_kw}", "kWe", "={gen_kva}", "kVA @ pf", "KBR max power × PowerMult (owner method)", "I.F Electrical Power max, Note 7", "H"),
    ("", "111-F", "Emergency Power Diesel Tank", "1", "=ROUNDUP({diesel_nom},0)", "m³ nominal", "={diesel_net}", "m³ net", "Backup h × genset kWe ÷ η ÷ LHV ÷ ρ", "I.F; IPCC LHV; EN 590", "H"),
    ("102 Raw Water", "121-F", "Raw Water Tank", "1", "=ROUNDUP({raw_nom},0)", "m³ nominal + fire reserve (TBD)", "={raw_net}", "m³ net", "Storage h × (CT make-up + demin feed + potable)", "Derived from I.F", "M"),
    ("", "121-JA/B", "Service Water Pumps", "2 (1+1)", "={raw}", "m³/h each", None, "", "Continuous raw water demand; + hose allowance TBD", "Derived from I.F", "M"),
    ("103 Demineralized Water", "131-L", "Demineralised Water Package", "1", "={demin_des}", "m³/h product", "={demin_feed}", "m³/h raw feed", "Demin max × margin", "I.F Demin (hub §3.3 for scaling)", "M"),
    ("", "131-F", "Demineralised Water Tank", "1", "=ROUNDUP({demin_nom},0)", "m³ nominal", "={demin_net}", "m³ net", "Storage h × design flow", "I.F", "M"),
    ("", "131-JA/B", "Demineralised Water Pumps", "2 (1+1)", "={demin_des}", "m³/h each", None, "", "Design flow", "I.F", "M"),
    ("", "132-F", "Neutralization Tank", "1", "=ROUNDUP({neut_nom},0)", "m³ nominal", "={neut_net}", "m³ net", "Storage h × regeneration waste", "A7", "L"),
    ("", "132-JA/B", "Waste Water Pumps", "2 (1+1)", "={demin_waste}", "m³/h each", None, "", "Regeneration waste (vendor minimum governs)", "A7", "L"),
    ("104 Polishing Water", "141-L", "Mixed Bed Polishing Package", "1", "={demin_des}", "m³/h", None, "", "Demin max × margin", "I.F", "M"),
    ("", "141-F", "Polished Water Tank", "1", "=ROUNDUP({demin_nom},0)", "m³ nominal", "={demin_net}", "m³ net", "Storage h × design flow", "I.F", "M"),
    ("", "141-JA/B", "Polished Water Pumps", "2 (1+1)", "={demin_des}", "m³/h each", "=KBR_DeminP", "barg at ISBL BL", "Design flow at BL pressure", "I.F", "M"),
    ("105 Potable Water", "151-F", "Potable Water Tank", "1", "=ROUNDUP({pot_nom},0)", "m³ nominal", "={pot_net}", "m³ net", "Headcount × L/p/d + safety showers", "WHO; ANSI Z358.1; A8", "L"),
    ("", "151-JA/B", "Potable Water Pumps", "1 in I.H (recommend 2)", "={pot_pump}", "m³/h each", None, "", "Safety showers + domestic peak", "ANSI Z358.1; A8", "L"),
    ("", "151-L", "Chlorination Package", "1", "={pot_avg}", "m³/h average fill", None, "", "Dose to ≥ 0.5 mg/L free Cl₂ after 30 min (WHO GDWQ)", "WHO GDWQ", "L"),
    ("106 Fire Water", "161-JA/B", "Fire Water Pumps (I.H: \"Service Water Pumps\")", "2", TBD, "", None, "", "Needs Fire & Explosion Risk Assessment + authority requirement", "No KBR data", "TBD"),
    ("107 Cooling Water System", "171-JA/B/C", "Cooling Water Pumps", "3 (2+1)", "={cw_each}", "m³/h each", "={cw_des}", "m³/h total", "CW max × margin; ≥ 5.0 barg discharge", "I.F C.W. max; hub §3.3", "H"),
    ("", "171-DA/B", "Cooling Water Towers", "2 cells", "={cw_duty}/2", "kW per cell", "={cw_each}", "m³/h per cell", "Q = ṁ cp ΔT, 30 → 40 °C", "I.F; I.A §6", "M"),
    ("", "171-L", "Chemical Injection Package", "1", "={ct_mu}/1000", "m³/h make-up", "={cw_des}", "m³/h circulation", "Dosing rates by water-treatment vendor", "Derived", "L"),
    ("108 Ammonia Cracking Flare / Fuel Gas", "181-L", "HP Flare Package", "1", "={flare_kg}", "kg/h", "={flare_nm3}", "Nm³/h", "PROXY: full cracked gas — replace with relief study", "I.E NH₃ feed", "L"),
    ("", "181-D", "HP Flare KO Drum", "1", "={flare_kg}", "kg/h gas", None, "", "Dimension per API Std 521 at FEED", "I.E", "L"),
    ("", "181-C", "HP Flare KO Drum Heater", "1", TBD, "", None, "", "Depends on liquid inventory; no KBR data", "—", "TBD"),
    ("", "182-L", "Fuel Gas Metering Package", "1", "={ng_des}", "kg/h", "={ng_nm3}", "Nm³/h", "NG max × margin at 7.5 barg", "I.E NG fuel; I.F ratio", "H"),
    ("", "182-D", "Fuel Gas KO Drum", "1", "={ng_des}", "kg/h", "={ng_nm3}", "Nm³/h", "As 182-L", "I.E; I.F", "H"),
    ("109 Air Systems", "191-JA/B/C", "Air Compressor Package A/B/C", "3 (2+1)", "={comp_each}", "Nm³/h each", "={comp_tot}", "Nm³/h total", "PA + IA dryer inlet, ≥ 8.5 barg", "I.F air max; A9", "M"),
    ("", "191-D", "Wet Air Receiver", "1", "={wet_rcv}", "m³", None, "", "Hold-up of compressor output", "A10", "L"),
    ("", "191-L", "Instrument Air Dryer Package", "1", "={ia_des}", "Nm³/h dry", "={dryer_in}", "Nm³/h inlet", "IA max × margin", "I.F IA max", "M"),
    ("", "192-D", "Instrument Air Receiver", "1", "={ia_rcv}", "m³", None, "", "IA hold-up 7.5 → 4.0 barg", "A10", "M"),
    ("110 Nitrogen System", "201-L", "Nitrogen Generation Package", "1", "={n2_gen}", "Nm³/h", "=KBR_N2P", "barg", "N₂ normal × margin (purity TBD)", "I.F N₂", "M"),
    ("", "201-D", "Liquid Nitrogen Storage Vessel", "1", "=ROUNDUP({ln2_nom},0)", "m³ nominal", "={ln2_net}", "m³ net LN₂", "(Start-up peak − generator) × duration", "I.F N₂; A11", "L"),
    ("", "201-CA/B", "Nitrogen Ambient Vaporisers A/B", "2 (2×100%)", "={vap_each}", "Nm³/h each", None, "", "N₂ max × margin", "I.F N₂ max", "M"),
    ("111 Off-Spec Tank", "211-F", "Off-Spec Tank", "1", TBD, "", None, "", "Service not defined by KBR", "—", "TBD"),
    ("112 Waste Water Treatment", "221-JA/B", "Neutralization Sump Pump A/B", "2 (1+1)", "={demin_waste}", "m³/h each", None, "", "Regeneration waste (may duplicate 132-JA/B)", "A7", "L"),
    ("", "221-L", "Oil Water Treatment Package", "1", TBD, "", None, "", "Needs rainfall intensity + paved area", "—", "TBD"),
    ("", "222-L", "Sanitary Lifting Station", "1", "={pot_day}", "m³/d", None, "", "Equal to potable use", "A8", "L"),
    ("", "222-JA/B", "Waste Water Effluent Sump Pump A/B", "2 (1+1)", "={effluent}", "m³/h + storm (TBD)", None, "", "Steam BD + CT BD + demin waste", "hub §3.5; derived", "M"),
    ("113 Ammonia Plant Hydrogen Export", "232-L", "Hydrogen Metering Package", "1", "={h2_des}", "kg/h", "={h2_nm3}", "Nm³/h", "H₂ × margin; turndown per I.A §9", "I.E H₂; I.C", "H"),
    ("", "231-L", "Hydrogen Pipeline Pig Launcher Package", "1", "={pipe_dn}", "indicative line", "={pipe_id}", "mm min ID", "Velocity ≤ PipeVmax", "A12; ASME B36.10M", "L"),
    ("114 Ammoniacal Water Drain System", "241-D", "Ammoniacal Drain Underground Drum", "1", TBD, "", None, "", "ISBL drain inventories not in KBR package", "—", "TBD"),
    ("", "241-JA/B", "Submerged Ammoniacal Drain Drum Pump", "2 (1+1)", TBD, "", None, "", "Follows 241-D", "—", "TBD"),
]

CONF_FILL = {"H": "C6E0B4", "M": "FFE699", "L": "FCE4D6", "TBD": "F8CBAD"}


def sub(formula, cells, col):
    out = formula.replace("{COL}", col)
    return re.sub(r"\{(\w+)\}", lambda m: cells[m.group(1)], out)


def build_cover(wb):
    ws = wb.active
    ws.title = "Cover"
    title(ws, "KBR H2ACT® — OSBL Equipment Sizing", "Gentari Hydrogen · Ammonia Cracker Workstream · Johor (Pasir Gudang)")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 90
    rows = [
        ("Scope", "Preliminary sizing of KBR's OSBL equipment list (I.H, Units 101–114) at 12, 24, 80 and 160 kTPA H₂."),
        ("Discipline", "Process / utilities engineering"),
        ("Version", f"{VERSION} — {DATE}"),
        ("Status", "DERIVED, PRELIMINARY — not a KBR deliverable. Basis for FEED/EPC discussion only."),
        ("Source data", "KBR Gentari proposal package Rev 0 (Dec 2025): I.A, I.C, I.E, I.F, I.H + H2ACT Technical Information Package. Indicative, proposal-stage."),
        ("Disclaimer", "KBR data is commercially confidential and indicative. I.F utility maxima exist for 12 kTPA only; other capacities are scaled (see Guide). "
                       "160 kTPA is ABOVE KBR's largest published case (80 kTPA) and is a linear extrapolation. No cost figures are carried in this workbook."),
        ("Standards cited", "ANSI/ISEA Z358.1 (safety showers); ASME B36.10M (pipe); API Std 521 (flare KO drum); ISO 8528 (genset rating); EN 590 (diesel); "
                            "IPCC 2006 GL (diesel LHV); WHO GDWQ (potable)"),
        ("Generator", "tools/kbr_osbl_workbook.py (regenerate; do not hand-edit formulas)"),
    ]
    for i, (k, v) in enumerate(rows, 5):
        style(ws.cell(row=i, column=1, value=k), bold=True, bg=F_CALC)
        style(ws.cell(row=i, column=2, value=v))
    r = 5 + len(rows) + 1
    ws.cell(row=r, column=1, value="Navigation").font = font(size=13, bold=True, color=F_HEAD)
    for j, sh in enumerate(["Guide", "Inputs", "KBR_Data", "Summary"] + [c[0] for c in CAPS], r + 1):
        c = ws.cell(row=j, column=1, value=f"▶ {sh}")
        c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{sh}'!A1")
        c.font = font(color="0563C1", underline="single")
    r = j + 2
    ws.cell(row=r, column=1, value="Revision history").font = font(size=13, bold=True, color=F_HEAD)
    header_row(ws, r + 1, ["Version / date", "Description"])
    style(ws.cell(row=r + 2, column=1, value=f"{VERSION} / {DATE}"))
    style(ws.cell(row=r + 2, column=2, value="First issue: 12 / 24 / 80 / 160 kTPA sheets, Summary, Inputs, KBR_Data. Extends tcoedatabase/KBR_OSBL_Sizing_12ktpa.md."))


def build_guide(wb):
    ws = wb.create_sheet("Guide")
    title(ws, "User Guide", "How the sizing works, what is assumed, and what is still open")
    ws.column_dimensions["A"].width = 130
    lines = [
        ("h", "1. How to use"),
        ("", "Change assumptions in the yellow cells on 'Inputs'; every sheet recalculates. KBR figures live on 'KBR_Data' (cited)."),
        ("", "No sheet is protected (owner request). Edit formulas on the capacity sheets with care: they are generated by tools/kbr_osbl_workbook.py."),
        ("", "Each capacity sheet has: (1) the sized equipment list, (2) the documented calculations beneath it (equation, formula, reference)."),
        ("", "Confidence: H = mostly KBR data · M = KBR data + material assumptions · L = mostly assumptions · TBD = cannot be sized from available data."),
        ("h", "2. Sizing rules"),
        ("", "Emergency power (owner method): KBR max ISBL power × 2, 24 h backup, 50% genset efficiency, diesel LHV. The ×2 is a LOAD ALLOWANCE, not redundancy: "
             "I.F Note 7 says the 540 kW covers ISBL process users only and excludes CT pumps/fans, N₂ generator, lighting, buildings and instruments. "
             "It is a single 1,080 kWe package (I.H: 111-L qty 1), not 2 × 540 kW. Compare with I.F Note 9 (~0.3 MW safe-shutdown estimate) and the "
             "estimated OSBL load in section M of each capacity sheet."),
        ("", "All other items: design flow = KBR maximum × 1.10 margin; tanks = 24 h autonomy ÷ 0.9 fill; A/B = 2 × 100%, A/B/C = 3 × 50%."),
        ("h", "3. How each capacity gets its KBR basis (important)"),
        ("", "KBR's I.F Utility Summary (maximum flows) is published for 12 kTPA ONLY. I.E Feed & Product gives power, NG, NH₃ and H₂ for 12 / 24 / 68 / 80 kTPA."),
        ("", "S1 — Power & NG: I.E normal value at each capacity × the I.F max/normal ratio at 12 kTPA (540/487 power; 450/418 NG). ASSUMPTION that the ratio holds."),
        ("", "S2 — Cooling water & demin: I.F 12 kTPA maximum × (capacity ÷ 12). Supported by KBR hub §3.3: CW 25 t/t H₂ and DNW 0.1 t/t H₂ are constant 12–80 kTPA."),
        ("", "S3 — Plant air, instrument air, nitrogen: I.F 12 kTPA × (capacity ÷ 12). ASSUMPTION — likely CONSERVATIVE, since air/N₂ demand does not scale linearly with throughput."),
        ("", "S4 — 160 kTPA: every capacity-specific KBR value = 80 kTPA × 2. ASSUMPTION — above KBR's largest published case; equivalent to 2 × 80 kTPA. "
             "KBR states a single-train ceiling of ~1,200 MTPD, so one train may still be possible; confirm with KBR."),
        ("", "Firing mode: power and NH₃ feed use the HIGHEST published mode per capacity (conservative); NG uses 100% NG mode (the only NG-heavy case). "
             "At 24 kTPA only the 100% NG mode is published."),
        ("h", "4. Changes vs. tcoedatabase/KBR_OSBL_Sizing_12ktpa.md (12 kTPA)"),
        ("", "Flare proxy now uses I.E NH₃ feed in the highest-feed mode (10,290 kg/h, 100% cracked-gas) for all capacities, instead of I.C NG-mode 9,104 kg/h "
             "→ 12 kTPA flare 11,319 kg/h (was ~10,000)."),
        ("", "H₂ metering now uses I.E H₂ 1,426 kg/h for all capacities (I.C HMB exists only at 12 kTPA: 1,435 kg/h) → 12 kTPA 1,569 kg/h (was 1,579)."),
        ("h", "5. KBR discrepancies (flagged, not reconciled)"),
        ("", "Emergency power: owner basis (540 × 2) vs. I.F Note 9 ~0.3 MW safe-shutdown estimate — shown as a sensitivity on each sheet."),
        ("", "CW supply temperature: I.F 25 °C vs. I.A/I.D 30 °C. 25 °C is below the design wet bulb (see KBR_Data) → infeasible for a cooling tower; 30 °C used."),
        ("", "CW pressures: I.F 5.0/3.5 barg; I.A 2.0 supply / 4.0 return (looks inverted); I.D 4.0 supply. 5.0 barg used."),
        ("", "Demin: I.F 245 kg/h normal vs. hub 0.1 t/t H₂ ≈ 143 kg/h. CW: I.F 30.1 t/h normal vs. hub 25 t/t ≈ 35.7 t/h. I.F maxima used."),
        ("", "NG pressure 7.5 (I.F) vs 7 barg (I.A). H₂ pressure 27 (I.E/I.C) vs 20 min (I.A) vs 28 barg (hub). H₂ 1,435 (I.C) vs 1,426 kg/h (I.E)."),
        ("", "I.H: 161-JA/B under Fire Water labelled 'Service Water Pumps'; 151-JA/B quantity 1 despite A/B tag; 132-JA/B and 221-JA/B may duplicate."),
        ("h", "6. Open items (TBD) — owner / KBR input needed"),
        ("", "Fire water rate & duration (FERA / authority) → 161-JA/B and the fire reserve in 121-F (Unit 106 has no tank; reserve likely shared in 121-F)."),
        ("", "Headcount & safety-shower philosophy → Unit 105, 222-L. Headcount is one input for all capacities — confirm per capacity."),
        ("", "Relief load summary → 181-L / 181-D / 181-C. Off-spec tank service → 211-F. ISBL ammoniacal drain inventories → 241-D / 241-JA/B."),
        ("", "Rainfall intensity & paved area → 221-L, 222-JA/B. N₂ start-up peak duration & purity → 201-D, 201-L. Genset heat rate → 111-F."),
        ("", "Equipment quantities are kept as listed in I.H at every capacity. At 80/160 kTPA more cooling-tower cells, compressors, etc. may be preferred — decide at FEED."),
    ]
    for i, (kind, text) in enumerate(lines, 5):
        c = ws.cell(row=i, column=1, value=text)
        if kind == "h":
            c.font = font(size=13, bold=True, color=F_HEAD)
        else:
            c.font = font()
            c.alignment = Alignment(wrap_text=True, vertical="top")


def build_inputs(wb):
    ws = wb.create_sheet("Inputs")
    title(ws, "Inputs", "Design assumptions — the only sheet users edit (yellow cells)")
    header_row(ws, 5, ["Name", "Parameter", "Value", "Unit", "Source / assumption"], [18, 52, 12, 14, 90])
    r = 6
    for item in INPUTS:
        if item[1] is None:
            c = ws.cell(row=r, column=1, value=item[0])
            c.font = font(size=12, bold=True, color=F_HEAD)
            r += 1
            continue
        name, label, val, unit, src, (lo, hi) = item
        for col, v in enumerate([name, label, val, unit, src], 1):
            c = ws.cell(row=r, column=col, value=v)
            style(c, bg=F_INPUT if col == 3 else None, color="0000FF" if col == 3 else "000000")
        vc = ws.cell(row=r, column=3)
        vc.protection = Protection(locked=False)
        dv = DataValidation(type="decimal", operator="between", formula1=str(lo), formula2=str(hi), showErrorMessage=True,
                            errorTitle="Invalid input", error=f"Enter a value between {lo} and {hi} {unit}.")
        ws.add_data_validation(dv)
        dv.add(vc)
        add_name(wb, name, f"Inputs!$C${r}")
        r += 1
    ws.freeze_panes = "A6"


def build_kbr(wb):
    ws = wb.create_sheet("KBR_Data")
    title(ws, "KBR Data", "KBR figures with document references (Licensor/kbr/, Rev 0, Dec 2025)")
    header_row(ws, 5, ["Key", "Parameter", "12 kTPA", "24 kTPA", "80 kTPA", "160 kTPA", "Unit", "KBR reference"],
               [14, 58, 12, 12, 12, 14, 10, 90])
    for i, (key, label, unit, vals, src) in enumerate(CAP_TABLE):
        r = 6 + i
        ws.cell(row=r, column=1, value=key)
        ws.cell(row=r, column=2, value=label)
        for j, v in enumerate(vals):
            ws.cell(row=r, column=3 + j, value=v)
        ws.cell(row=r, column=6, value=160 if key == "cap" else f"=E{r}*$F$6/$E$6")
        ws.cell(row=r, column=7, value=unit)
        ws.cell(row=r, column=8, value=src)
        for col in range(1, 9):
            c = ws.cell(row=r, column=col)
            style(c, bg=F_EXTRAP if col == 6 else None, fmt=NUM if 3 <= col <= 6 else None)
    ws.cell(row=6, column=6).comment = Comment("160 kTPA column = 80 kTPA × 2 (ASSUMPTION S4): above KBR's largest published case.", "Gentari TCOE")
    r = 6 + len(CAP_TABLE)
    ws.cell(row=r, column=2, value="Orange column: 160 kTPA = 80 kTPA × (160/80) — linear extrapolation, ASSUMPTION S4.").font = font(italic=True, color="C55A11")
    r += 2
    header_row(ws, r, ["Named range", "Parameter", "Value", "", "", "", "Unit", "KBR reference"])
    r += 1
    for item in KBR_SINGLE:
        if item[1] is None:
            ws.cell(row=r, column=1, value=item[0]).font = font(size=12, bold=True, color=F_HEAD)
            r += 1
            continue
        name, label, val, unit, src = item
        for col, v in [(1, name), (2, label), (3, val), (7, unit), (8, src)]:
            style(ws.cell(row=r, column=col, value=v), fmt=NUM if col == 3 else None)
        add_name(wb, name, f"KBR_Data!$C${r}")
        r += 1
    # wet bulb check (Stull 2011)
    r += 1
    ws.cell(row=r, column=1, value="Check: CW supply temperature vs design wet bulb").font = font(size=12, bold=True, color=F_HEAD)
    r += 1
    t, rh = "KBR_AmbT", "KBR_AmbRH"
    wb_f = (f"={t}*ATAN(0.151977*({rh}+8.313659)^0.5)+ATAN({t}+{rh})-ATAN({rh}-1.676331)"
            f"+0.00391838*{rh}^1.5*ATAN(0.023101*{rh})-4.686035")
    checks = [
        ("Design wet bulb (Stull 2011) at I.A ambient", wb_f, "°C", "Stull, J. Appl. Meteor. Climatol. 50 (2011) 2267"),
        ("Approach with I.F 25 °C supply", f"=25-C{r}", "°C", "Negative → infeasible for a cooling tower"),
        ("Approach with I.A/I.D 30 °C supply (used)", f"=KBR_CWSupplyT-C{r}", "°C", "Positive → feasible (use 1% exceedance wet bulb at FEED)"),
    ]
    for k, (lab, f, unit, src) in enumerate(checks):
        rr = r + k
        for col, v in [(2, lab), (3, f), (7, unit), (8, src)]:
            style(ws.cell(row=rr, column=col, value=v), bg=F_CALC if col == 3 else None, fmt="0.0" if col == 3 else None)
    r += len(checks) + 1
    ws.cell(row=r, column=1, value="Pipe table — ASME B36.10M Schedule 40 inside diameters").font = font(size=12, bold=True, color=F_HEAD)
    r += 1
    header_row(ws, r, ["", "Nominal size", "ID (mm)", "", "", "", "", ""])
    first = r + 1
    for k, (dn, idmm) in enumerate(PIPES):
        style(ws.cell(row=first + k, column=2, value=dn))
        style(ws.cell(row=first + k, column=3, value=idmm), fmt="0.0")
    last = first + len(PIPES) - 1
    add_name(wb, "PipeDN", f"KBR_Data!$B${first}:$B${last}")
    add_name(wb, "PipeID", f"KBR_Data!$C${first}:$C${last}")
    ws.freeze_panes = "C6"


def build_capacity(wb, sheet, cap):
    col = CAP_COLS[cap]
    ws = wb.create_sheet(sheet)
    extrap = cap > 80
    title(ws, f"OSBL Equipment Sizing — {sheet} H₂",
          "KBR H2ACT® · Pasir Gudang, Johor · derived from KBR Rev 0 data" + (" · EXTRAPOLATED BEYOND KBR's 80 kTPA CASE" if extrap else ""))
    if extrap:
        ws["A2"].font = font(italic=True, bold=True, color="C55A11")
    widths = [30, 12, 36, 14, 14, 22, 14, 18, 50, 30, 8]
    header_row(ws, 5, ["Unit", "Tag", "Item", "Qty", "Design capacity", "Unit", "Secondary", "Unit", "Basis", "KBR / source ref", "Conf."], widths)

    # lay out calculations below the equipment list, record cell addresses first
    calc_start = 5 + 1 + len(EQUIP) + 3
    cells, r = {}, calc_start + 1
    for item in CALCS:
        if item[1] is not None:
            cells[item[0]] = f"$C${r}"
        r += 1

    # equipment list
    for i, (unit, tag, name, qty, f1, u1, f2, u2, basis, ref, conf) in enumerate(EQUIP):
        rr = 6 + i
        v1 = f1 if f1 == TBD else sub(f1, cells, col)
        v2 = sub(f2, cells, col) if f2 else None
        vals = [unit, tag, name, qty, v1, u1, v2, u2, basis, ref, conf]
        for c_i, v in enumerate(vals, 1):
            c = ws.cell(row=rr, column=c_i, value=v)
            bg = None
            if c_i in (5, 7) and v is not None:
                bg = F_TBD if v == TBD else F_CALC
            if c_i == 11:
                bg = CONF_FILL[conf]
            style(c, bg=bg, bold=(c_i in (1, 5)), fmt=NUM if c_i in (5, 7) else None,
                  align="right" if c_i in (5, 7) else "left")

    # calculations
    ws.cell(row=calc_start - 1, column=1, value="Documented calculations").font = font(size=14, bold=True, color=F_HEAD)
    hdr = ["Key", "Parameter", "Value", "Unit", "Equation / reference", "", "", "", "Excel formula", "", ""]
    for i, lab in enumerate(hdr, 1):
        style(ws.cell(row=calc_start, column=i, value=lab), bg=F_HEAD, bold=True, color="FFFFFF")
    r = calc_start + 1
    for item in CALCS:
        if item[1] is None:
            c = ws.cell(row=r, column=1, value=item[0])
            c.font = font(size=12, bold=True, color=F_HEAD)
            r += 1
            continue
        key, label, f, unit, eq = item
        formula = sub(f, cells, col)
        style(ws.cell(row=r, column=1, value=key), color="808080")
        style(ws.cell(row=r, column=2, value=label))
        style(ws.cell(row=r, column=3, value=formula), bg=F_EXTRAP if (extrap and key in ("cap", "p_norm", "ng_norm", "nh3_max", "h2")) else F_CALC,
              fmt=NUM, align="right")
        style(ws.cell(row=r, column=4, value=unit))
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
        style(ws.cell(row=r, column=5, value=eq))
        ws.merge_cells(start_row=r, start_column=9, end_row=r, end_column=11)
        style(ws.cell(row=r, column=9, value=formula.lstrip("=")), color="808080")
        r += 1
    ws.freeze_panes = "D6"
    return cells


def build_summary(wb, cells_by_sheet):
    ws = wb.create_sheet("Summary", 4)
    title(ws, "Summary — OSBL Design Capacity by H₂ Capacity", "Primary design capacity per item; see each capacity sheet for secondary values and calculations")
    header_row(ws, 5, ["Unit", "Tag", "Item", "Unit of capacity"] + [c[0] for c in CAPS] + ["Conf."],
               [30, 12, 36, 26, 14, 14, 14, 16, 8])
    for i, (unit, tag, name, qty, f1, u1, f2, u2, basis, ref, conf) in enumerate(EQUIP):
        rr = 6 + i
        vals = [unit, tag, name, u1]
        for c_i, v in enumerate(vals, 1):
            style(ws.cell(row=rr, column=c_i, value=v), bold=c_i == 1)
        for j, (sheet, _) in enumerate(CAPS):
            c = ws.cell(row=rr, column=5 + j, value=TBD if f1 == TBD else f"='{sheet}'!E{rr}")
            style(c, bg=F_TBD if f1 == TBD else (F_EXTRAP if j == 3 else F_CALC), fmt=NUM, align="right")
        style(ws.cell(row=rr, column=9, value=conf), bg=CONF_FILL[conf])
    r = 6 + len(EQUIP) + 2
    ws.cell(row=r, column=1, value="Key utility basis").font = font(size=13, bold=True, color=F_HEAD)
    r += 1
    header_row(ws, r, ["Parameter", "", "", "Unit"] + [c[0] for c in CAPS] + [""])
    keys = [("p_max", "Electrical power, maximum (KBR basis)", "kW"), ("gen_kw", "Emergency genset rating", "kWe"),
            ("cw_des", "Cooling water circulation, design", "m³/h"), ("cw_duty", "Cooling water duty", "kW"),
            ("raw", "Raw water, continuous", "m³/h"), ("ng_des", "Fuel gas, design", "kg/h"),
            ("h2_des", "H₂ export, design", "kg/h"), ("allow", "Power allowance from multiplier", "kW"),
            ("osbl_est", "Estimated OSBL rotating load (partial)", "kW")]
    for k, (key, lab, unit) in enumerate(keys, r + 1):
        style(ws.cell(row=k, column=1, value=lab), bold=True)
        ws.merge_cells(start_row=k, start_column=1, end_row=k, end_column=3)
        style(ws.cell(row=k, column=4, value=unit))
        for j, (sheet, _) in enumerate(CAPS):
            addr = cells_by_sheet[sheet][key].replace("$", "")
            style(ws.cell(row=k, column=5 + j, value=f"='{sheet}'!{addr}"), bg=F_EXTRAP if j == 3 else F_CALC, fmt=NUM, align="right")
    k += 2
    ws.cell(row=k, column=1, value="Orange = 160 kTPA, extrapolated beyond KBR's 80 kTPA case (80 × 2). Red = TBD. "
                                   "Confidence: H/M/L as defined on Guide.").font = font(italic=True, color="595959")
    ws.freeze_panes = "E6"


def main():
    wb = Workbook()
    build_cover(wb)
    build_guide(wb)
    build_inputs(wb)
    build_kbr(wb)
    cells_by_sheet = {sheet: build_capacity(wb, sheet, cap) for sheet, cap in CAPS}
    build_summary(wb, cells_by_sheet)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
